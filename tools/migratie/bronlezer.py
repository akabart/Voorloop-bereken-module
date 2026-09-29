"""Hulpfuncties voor het inlezen van de bronbestanden (xls, xlsx, pdf) en het normaliseren van waarden."""
import re

import openpyxl
import xlrd


# ---------------------------------------------------------------------------
# Werkbladen: één interface voor .xls en .xlsx
# ---------------------------------------------------------------------------

class Blad:
    """Werkblad met 0-gebaseerde toegang: blad.v(rij, kolom) geeft de (gecachte) waarde of None."""

    def __init__(self, naam, cellen, nrows, ncols, verborgen, samengevoegd=()):
        self.samengevoegd = list(samengevoegd)   # (r1, r2, c1, c2) met exclusieve einden
        self.naam = naam
        self._cellen = cellen
        self.nrows = nrows
        self.ncols = ncols
        self.verborgen = verborgen

    def v(self, r, c):
        val = self._cellen.get((r, c))
        if isinstance(val, str):
            val = val.replace('\xa0', ' ').strip()
            return val or None
        return val

    def kop(self, r, c):
        """Waarde van een (kop)cel; bij een samengevoegde cel de waarde linksboven."""
        v = self.v(r, c)
        if v is not None:
            return v
        for r1, r2, c1, c2 in self.samengevoegd:
            if r1 <= r < r2 and c1 <= c < c2:
                return self.v(r1, c1)
        return None

    def rij(self, r):
        return [self.v(r, c) for c in range(self.ncols)]


def open_werkboek(pad):
    if pad.lower().endswith('.xls'):
        boek = xlrd.open_workbook(pad, formatting_info=True)
        bladen = []
        for sh in boek.sheets():
            cellen = {}
            for r in range(sh.nrows):
                for c in range(sh.ncols):
                    cel = sh.cell(r, c)
                    if cel.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                        continue
                    cellen[(r, c)] = cel.value
            bladen.append(Blad(sh.name, cellen, sh.nrows, sh.ncols, sh.visibility != 0, sh.merged_cells))
        return bladen
    boek = openpyxl.load_workbook(pad, data_only=True)
    bladen = []
    for ws in boek.worksheets:
        cellen = {}
        for row in ws.iter_rows():
            for cel in row:
                if cel.value is not None:
                    cellen[(cel.row - 1, cel.column - 1)] = cel.value
        samen = [(m.min_row - 1, m.max_row, m.min_col - 1, m.max_col) for m in ws.merged_cells.ranges]
        bladen.append(Blad(ws.title, cellen, ws.max_row, ws.max_column, ws.sheet_state != 'visible', samen))
    return bladen


def kolomnaam(c):
    s = ''
    c += 1
    while c:
        c, rest = divmod(c - 1, 26)
        s = chr(65 + rest) + s
    return s


# ---------------------------------------------------------------------------
# Normalisatie
# ---------------------------------------------------------------------------

GETAL = re.compile(r'\d+(?:[.,]\d+)?')


def tekst(v):
    """Waarde als nette tekst (getallen zonder overbodige .0)."""
    if v is None:
        return None
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return ('%.6f' % v).rstrip('0').rstrip('.').replace('.', ',')
    s = str(v).strip()
    return s or None


def getal(v):
    """Eerste getal uit een waarde (komma of punt als decimaalteken), anders None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = GETAL.search(str(v))
    return float(m.group(0).replace(',', '.')) if m else None


LEEG = {'----', '---', '--', '-', 'n/a', 'xxx', ',nn'}


# Een verhouding ziet eruit als 1,xxx (of 0,xxx in Fendt-notatie), of als 1xxx zonder decimaalteken.
RATIO = re.compile(r'(?<![\d.,])(?:[01][.,]\d{2,6}|1\d{3,4})(?![\d])')
TOEGESTANE_WOORDEN = re.compile(r'\b(std|or|of|if|cvt|serv|kl\s*\d|va|aa|en)\b', re.I)


def verhoudingen(ruw, notatie='direct'):
    """Zet een ruwe celwaarde om naar genormaliseerde verhoudingen i = n_voor / n_achter.

    notatie 'direct': waarde is al i (> 1).
    notatie 'va_ha' : Fendt-notatie (i_VA / i_HA < 1), i = 1 / waarde.
    Geeft (lijst van waarden, lijst van opmerkingen). Een cel zonder herkenbare verhouding geeft ([], []).
    """
    if ruw is None:
        return [], []
    opm = []
    if isinstance(ruw, (int, float)):
        s = tekst(ruw)
        kandidaten = [float(ruw)]
        if kandidaten[0] >= 10 and not RATIO.fullmatch(str(int(kandidaten[0]))):
            return [], []
    else:
        s = str(ruw).strip()
        if s.lower() in LEEG or 'niet leverbaar' in s.lower():
            return [], []
        kandidaten = [float(x.replace(',', '.')) for x in RATIO.findall(s)]
        if not kandidaten:
            return [], []
        rest = TOEGESTANE_WOORDEN.sub('', RATIO.sub('', s))
        if re.search(r'[a-zA-Z]', rest):
            opm.append('tekst in de cel: "%s"' % s)
    uit = []
    for v in kandidaten:
        if v <= 0:
            continue
        if v >= 10:
            orig = v
            while v >= 10:
                v /= 10
            opm.append('decimaalteken ontbrak (%s gelezen als %s)' % (tekst(orig), ('%.4f' % v).replace('.', ',')))
        if notatie == 'va_ha':
            if v > 1.0:
                opm.append('Fendt-notatie verwacht (< 1), maar %s gevonden' % tekst(v))
            else:
                v = 1.0 / v
        elif v < 1.0:
            opm.append('waarde %s < 1: omgekeerd gelezen als %s' % (tekst(v), ('%.4f' % (1 / v)).replace('.', ',')))
            v = 1.0 / v
        if not (1.0 <= v <= 1.8):
            opm.append('buiten plausibel bereik: %s' % ('%.4f' % v).replace('.', ','))
            continue
        uit.append(round(v, 5))
    if len(uit) > 1:
        opm.append('meerdere waarden in één cel: "%s"' % s)
    return uit, opm


def soort_probleem(opm):
    """Categorie voor het reviewscherm (zodat eenvoudige gevallen in één keer goed te keuren zijn)."""
    tekst_ = ' '.join(opm)
    if 'meerdere waarden' in tekst_:
        return 'meerdere'
    if 'buiten plausibel' in tekst_ or '< 1' in tekst_ or 'Fendt-notatie' in tekst_ or '0 in bron' in tekst_:
        return 'plausibiliteit'
    if 'tekst in de cel' in tekst_ or 'PDF' in tekst_:
        return 'tekst'
    if opm and all('decimaalteken ontbrak' in o for o in opm):
        return 'decimaal'
    return 'overig'


def bouten_draad(ruw):
    """'8/16X1.5' -> (8, 'M16x1,5'); '10/3/4' -> (10, '3/4"'); '8/M22x1.5 220nm' -> (8, 'M22x1,5')."""
    if ruw is None:
        return None, None
    s = tekst(ruw)
    s = re.sub(r'\s*\d+\s*nm\b', '', s, flags=re.I).strip()
    m = re.match(r'^\s*(\d+)\s*(?:[/x×]|\s)\s*(.+)$', s, flags=re.I)
    if not m:
        return None, draad(s)
    return int(m.group(1)), draad(m.group(2))


def draad(s):
    if not s:
        return None
    s = s.strip().replace('X', 'x').replace(' ', '')
    if re.match(r'^\d+x\d', s):
        s = 'M' + s
    s = re.sub(r'^m', 'M', s)
    s = s.replace('.', ',')
    if re.match(r'^\d+/\d+$', s):
        s += '"'
    return s


def aanhaalmoment(ruw):
    if ruw is None:
        return None
    m = re.search(r'(\d+)\s*nm', str(ruw), flags=re.I)
    return int(m.group(1)) if m else None


def schoon_mm(v):
    """'1920 mm' -> '1920'; 150.5 -> '150,5'; laat tekst als '1890/1900' intact."""
    if v is None:
        return None
    s = tekst(v)
    s = re.sub(r'\s*mm\b', '', s, flags=re.I).strip()
    s = s.replace('.', ',')
    return s or None
